from typing import List, Dict, Any
from pathlib import Path
from pptx import Presentation

class PPTXParser:
    """
    Extracts text, slide titles, bullet points, tables, and speaker notes
    from PowerPoint presentations while preserving exact 1-indexed slide numbers.
    """
    
    @staticmethod
    def extract_slides(file_path: str) -> List[Dict[str, Any]]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"PPTX file not found: {file_path}")
        
        extracted_slides = []
        try:
            prs = Presentation(str(path))
            for slide_idx, slide in enumerate(prs.slides):
                slide_num = slide_idx + 1  # 1-indexed
                title = ""
                content_lines = []
                notes_text = ""
                
                # Check slide title
                if slide.shapes.title and slide.shapes.title.text:
                    title = slide.shapes.title.text.strip()
                
                for shape in slide.shapes:
                    if shape.has_text_frame and shape != slide.shapes.title:
                        for paragraph in shape.text_frame.paragraphs:
                            p_text = paragraph.text.strip()
                            if p_text:
                                content_lines.append(p_text)
                    
                    # Extract tables if present
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_cells:
                                content_lines.append(" | ".join(row_cells))
                
                # If title not found via layout shape, take the first non-empty line as title
                if not title and content_lines:
                    first_cand = content_lines[0].lstrip("•- \t")
                    if len(first_cand) < 90 and not first_cand.endswith("."):
                        title = first_cand
                
                # Extract speaker notes
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
                
                full_text_parts = []
                if title:
                    full_text_parts.append(f"Title: {title}")
                if content_lines:
                    full_text_parts.append("\n".join(content_lines))
                if notes_text:
                    full_text_parts.append(f"Speaker Notes: {notes_text}")
                
                slide_text = "\n\n".join(full_text_parts).strip()
                headings = [title] if title else []
                
                extracted_slides.append({
                    "slide_number": slide_num,
                    "title": title,
                    "text": slide_text,
                    "notes": notes_text,
                    "headings": headings,
                    "extraction_method": "text",
                    "confidence": 0.98
                })
                
            return extracted_slides
        except Exception as e:
            raise RuntimeError(f"Error parsing PPTX file {file_path}: {str(e)}")
